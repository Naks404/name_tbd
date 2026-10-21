from __future__ import annotations
from copy import copy
from pathlib import Path
import cv2
import numpy as np
import time
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIG
# ============================================================
BIG_FOLDER = Path(r"C:\Users\Kyros Lialios\Downloads\IEA\IEA_img")
OUTPUT_XLSX = "exposure_stochastic_benchmark.xlsx"

# The sequence as defined: N4 < N3 < N2 < N1 < Ref < P0 < P1 < P2 < P3 < P4
SUFF_ORDER = ["_N4", "_N3", "_N2", "_N1", "", "_P0", "_P1", "_P2", "_P3", "_P4"]
OFFSET_MAP = {suffix: i - 4 for i, suffix in enumerate(SUFF_ORDER)}  # Ref (index 4) becomes 0


# ============================================================
# OPTIMIZED ALGORITHMS
# ============================================================
def prep_image(image):
    """Calculates grayscale and histogram once to save time."""
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    hist = cv2.calcHist([gray], [0], None, [256], [0, 256]).flatten()
    total_pixels = gray.shape[0] * gray.shape[1]
    return gray, hist, total_pixels


def mean_opt(gray, hist, total_pixels):
    return (cv2.mean(gray)[0],)


def median_opt(gray, hist, total_pixels):
    cumulative_hist = np.cumsum(hist)
    median = np.searchsorted(cumulative_hist, total_pixels / 2)
    return (median,)


def mode_opt(gray, hist, total_pixels):
    return (np.argmax(hist),)


def coverage_opt(gray, hist, total_pixels):
    return ((np.count_nonzero(hist) / 256.0) * 100,)


def outliers_opt(gray, hist, total_pixels):
    # Summing histogram bins is faster than np.count_nonzero on the full image
    blacks = np.sum(hist[:5])
    whites = np.sum(hist[251:])
    return (blacks, whites)


def entropy_opt(gray, hist, total_pixels):
    probs = hist / total_pixels
    probs = probs[probs > 0]
    return (-np.sum(probs * np.log2(probs)),)


def skewness_opt(gray, hist, total_pixels):
    probs = hist / total_pixels
    bins = np.arange(256)
    skewness = np.sum(((bins - 128) ** 3) * probs)
    return (skewness,)


def grad_opt(gray, hist, total_pixels):
    # Grad remains the bottleneck; we use float32 for a slight speedup over float64
    gray_norm = gray.astype(np.float32) / 255.0
    gx = cv2.Sobel(gray_norm, cv2.CV_32F, 1, 0, ksize=3)
    gy = cv2.Sobel(gray_norm, cv2.CV_32F, 0, 1, ksize=3)
    mi = cv2.magnitude(gx, gy) / np.sqrt(20.0)

    lamda, delta = 1000, 0.11
    N = np.log2(lamda * (1 - delta) + 1)

    mask = mi >= delta
    m_bar = np.zeros_like(mi)
    m_bar[mask] = (1 / N) * np.log2(lamda * (mi[mask] - delta) + 1)
    return (np.sum(m_bar),)


FUNCTIONS = {
    "mean": mean_opt, "median": median_opt, "mode": mode_opt,
    "coverage": coverage_opt, "outliers": outliers_opt,
    "entropy": entropy_opt, "skewness": skewness_opt, "grad": grad_opt,
}

RULES = {
    "mean": "closest_128", "median": "closest_128", "mode": "closest_128",
    "coverage": "max", "outliers": "lowest_sum_first_two",
    "entropy": "max", "skewness": "closest_0", "grad": "max",
}

# Excel Columns Mapping
SHEET1_COLUMNS = [
    ("mean1", "mean", 0), ("median1", "median", 0), ("mode1", "mode", 0),
    ("coverage1", "coverage", 0), ("outliers1", "outliers", 0), ("outliers2", "outliers", 1),
    ("entropy1", "entropy", 0), ("skewness1", "skewness", 0), ("grad1", "grad", 0),
]


# ============================================================
# HELPERS
# ============================================================
def get_selected_indices(rule: str, values_for_group):
    if rule == "closest_128":
        dists = [abs(float(v) - 128.0) for v in values_for_group]
        best = min(dists)
        return {i for i, d in enumerate(dists) if d == best}
    if rule == "closest_0":
        dists = [abs(float(v) - 0.0) for v in values_for_group]
        best = min(dists)
        return {i for i, d in enumerate(dists) if d == best}
    if rule == "max":
        best = max(float(v) for v in values_for_group)
        return {i for i, v in enumerate(values_for_group) if float(v) == best}
    if rule == "lowest_sum_first_two":
        sums = [float(v[0]) + float(v[1]) for v in values_for_group]
        best = min(sums)
        return {i for i, s in enumerate(sums) if s == best}
    return set()


def style_header(ws, fill, font):
    for cell in ws[1]:
        cell.fill, cell.font = fill, font
        cell.alignment = Alignment(horizontal="center")


# ============================================================
# MAIN EXECUTION
# ============================================================
def main():
    print("Mapping directory...")
    # Map aXYZW -> { suffix: path }
    file_map = {}
    pattern = re.compile(r"^(a\d{4})")
    for p in BIG_FOLDER.glob("*.png"):
        match = pattern.match(p.name)
        if match:
            group_id = match.group(1)
            if group_id not in file_map: file_map[group_id] = {}

            # Identify suffix
            found_suffix = ""
            for s in [s for s in SUFF_ORDER if s != ""]:
                if p.stem.endswith(s):
                    found_suffix = s
                    break
            file_map[group_id][found_suffix] = p

    wb = Workbook()
    ws1, ws2, ws3 = wb.active, wb.create_sheet("times"), wb.create_sheet("stochastic")
    ws1.title = "results"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    ws1.append(["Group", "Exposure", "Path"] + [c[0] for c in SHEET1_COLUMNS])
    ws2.append(["Group", "Exposure"] + list(FUNCTIONS.keys()))
    ws3.append(["Group"] + list(FUNCTIONS.keys()))
    for s in [ws1, ws2, ws3]: style_header(s, header_fill, header_font)

    hit_counts = {fn: 0 for fn in RULES}
    stochastic_data = {fn: [] for fn in RULES}

    sorted_group_ids = sorted(file_map.keys())

    for gid in sorted_group_ids:
        group_paths = file_map[gid]
        group_results = []
        group_times = []

        # Process the 10 versions in the specific N4 -> P4 order
        for suffix in SUFF_ORDER:
            path = group_paths.get(suffix)
            if not path:
                group_results.append(None)
                group_times.append(None)
                continue

            img = cv2.imread(str(path))
            # The Preprocessing Step
            gray, hist, total_px = prep_image(img)

            res, times = {}, {}
            for fname, func in FUNCTIONS.items():
                t0 = time.perf_counter()
                res[fname] = func(gray, hist, total_px)
                times[fname] = time.perf_counter() - t0

            group_results.append(res)
            group_times.append(times)

        # Analysis
        stoch_row = [gid]
        for fname, rule in RULES.items():
            valid_indices = [i for i, r in enumerate(group_results) if r is not None]
            vals = [group_results[i][fname] if fname == "outliers" else group_results[i][fname][0] for i in
                    valid_indices]

            winners = get_selected_indices(rule, vals)

            # Sheet 1 Hits: Index 4 is the Reference in the sorted SUFF_ORDER
            if 4 in winners: hit_counts[fname] += 1

            # Stochastic offset: Mean of winner indices mapped to -4/+5
            offsets = [OFFSET_MAP[SUFF_ORDER[i]] for i in winners]
            avg_offset = sum(offsets) / len(offsets)
            stochastic_data[fname].append(avg_offset)
            stoch_row.append(avg_offset)

        ws3.append(stoch_row)

        # Write Rows for Sheet 1 and 2
        for i, suffix in enumerate(SUFF_ORDER):
            if group_results[i] is None: continue

            # Sheet 1
            row1 = [gid, suffix if suffix else "REF", str(group_paths[suffix])]
            for _, fname, out_idx in SHEET1_COLUMNS:
                row1.append(group_results[i][fname][out_idx])
            ws1.append(row1)

            # Bold wins
            curr_row = ws1.max_row
            for col_idx, (_, fname, _) in enumerate(SHEET1_COLUMNS, start=4):
                winners = get_selected_indices(RULES[fname], [
                    group_results[j][fname] if fname == "outliers" else group_results[j][fname][0] for j in range(10) if
                    group_results[j]])
                if i in winners:
                    ws1.cell(row=curr_row, column=col_idx).font = Font(bold=True)

            # Sheet 2 (Times)
            ws2.append([gid, suffix if suffix else "REF"] + [group_times[i][f] for f in FUNCTIONS])

    # Final Summaries
    ws1.append([]);
    ws1.append(["Method", "Hits", "Total", "Hit Rate %"])
    for fn in FUNCTIONS:
        ws1.append([fn, hit_counts[fn], len(sorted_group_ids), (hit_counts[fn] / len(sorted_group_ids)) * 100])

    ws3.append([]);
    ws3.append(["Metric"] + list(FUNCTIONS.keys()))
    ws3.append(["Mean Error"] + [np.mean(stochastic_data[f]) for f in FUNCTIONS])
    ws3.append(["Median Error"] + [np.median(stochastic_data[f]) for f in FUNCTIONS])
    ws3.append(["Variance"] + [np.var(stochastic_data[f]) for f in FUNCTIONS])

    wb.save(OUTPUT_XLSX)
    print(f"Benchmark Complete: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()